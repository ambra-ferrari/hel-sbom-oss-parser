import sys
from pathlib import Path
import json as json_module
import openpyxl as _openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
OSS_TEMPLATE = str(Path(__file__).resolve().parent.parent / "templates" / "oss-template.xlsx")

import sbom_to_oss_dependencies as oss_dependencies
from sbom_to_oss_dependencies import (
    build_component_graph,
    build_excel,
    build_components,
    build_rows,
    resolve_license,
    resolve_reference,
    resolve_purl,
    resolve_purpose,
)


def test_first_party_dependency_config_prefers_source_tree(monkeypatch, tmp_path):
    source_config = tmp_path / "config" / "oss_deps_first_party_dependencies.json"
    source_config.parent.mkdir()
    source_config.write_text("{}")
    monkeypatch.setattr(
        oss_dependencies, "DEFAULT_OSS_DEPS_FIRST_PARTY_DEPENDENCIES", source_config
    )
    monkeypatch.setattr(
        oss_dependencies.sysconfig, "get_path",
        lambda scheme: str(tmp_path / "install-data"),
    )

    assert oss_dependencies.resolve_first_party_dependency_config_path() == source_config


def test_first_party_dependency_config_falls_back_to_packaged_data(monkeypatch, tmp_path):
    source_config = tmp_path / "source" / "config" / "oss_deps_first_party_dependencies.json"
    packaged_config = (
        tmp_path / "install-data" / "share" / "hel-sbom-oss-parser"
        / "oss_deps_first_party_dependencies.json"
    )
    packaged_config.parent.mkdir(parents=True)
    packaged_config.write_text("{}")
    monkeypatch.setattr(
        oss_dependencies, "DEFAULT_OSS_DEPS_FIRST_PARTY_DEPENDENCIES", source_config
    )
    monkeypatch.setattr(
        oss_dependencies.sysconfig, "get_path",
        lambda scheme: str(tmp_path / "install-data"),
    )

    assert oss_dependencies.resolve_first_party_dependency_config_path() == packaged_config


def test_first_party_dependency_config_uses_distribution_locate_file(
        monkeypatch, tmp_path):
    source_config = tmp_path / "source" / "config" / "oss_deps_first_party_dependencies.json"
    custom_target = tmp_path / "custom-prefix" / "share" / "hel-sbom-oss-parser"
    packaged_config = custom_target / "oss_deps_first_party_dependencies.json"
    custom_target.mkdir(parents=True)
    packaged_config.write_text("{}")
    monkeypatch.setattr(
        oss_dependencies, "DEFAULT_OSS_DEPS_FIRST_PARTY_DEPENDENCIES", source_config
    )

    class InstalledDistribution:
        files = [Path("custom-prefix/share/hel-sbom-oss-parser/"
                      "oss_deps_first_party_dependencies.json")]

        def locate_file(self, package_path):
            return tmp_path / package_path

    monkeypatch.setattr(
        oss_dependencies.metadata, "distribution",
        lambda name: InstalledDistribution(),
    )

    assert oss_dependencies.resolve_first_party_dependency_config_path() == packaged_config


def test_missing_first_party_dependency_config_fails_clearly(tmp_path):
    missing_config = tmp_path / "missing.json"

    with pytest.raises(FileNotFoundError, match=str(missing_config)):
        oss_dependencies.resolve_first_party_dependency_config_path(str(missing_config))


def test_first_party_dependency_config_rejects_non_object_json(tmp_path):
    config_path = tmp_path / "not-an-object.json"
    config_path.write_text("[]")

    with pytest.raises(ValueError, match="must contain a JSON object"):
        oss_dependencies.load_first_party_dependency_ids(str(config_path))


def test_first_party_dependency_config_rejects_invalid_json(tmp_path):
    config_path = tmp_path / "invalid.json"
    config_path.write_text("{")

    with pytest.raises(ValueError, match="not valid JSON"):
        oss_dependencies.load_first_party_dependency_ids(str(config_path))


def test_build_excel_accepts_first_party_dependency_config_path(tmp_path):
    packages = [
        {"SPDXID": "SPDXRef-internal", "name": "biz.videomed.tools:internal-helper",
         "versionInfo": "1.0", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-app", "name": "hel-app",
         "versionInfo": "2.0", "supplier": "Organization: Baxter"},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-internal", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-app"},
    ]
    config_path = tmp_path / "first-party-dependencies.json"
    config_path.write_text(json_module.dumps({"INTERNAL-HELPER": "test config"}))
    sbom = _write_sbom(tmp_path, packages, relationships)
    output = tmp_path / "out.xlsx"

    build_excel(
        sbom, OSS_TEMPLATE, str(output),
        first_party_dependency_config_path=str(config_path),
    )

    worksheet = _openpyxl.load_workbook(output)["Dependencies (OTS SOUP)"]
    assert (worksheet.cell(2, 1).value, worksheet.cell(2, 9).value) == (
        "biz.videomed.tools:internal-helper", "hel-app"
    )


def _write_sbom(tmp_path, packages, relationships=None):
    doc = {
        "spdxVersion": "SPDX-2.3",
        "name": "Test Product",
        "creationInfo": {"created": "2026-09-29T00:00:00Z",
                          "creators": ["Organization: Helion", "Tool: fossa-cli"]},
        "documentDescribes": [],
        "packages": packages,
        "relationships": relationships or [],
    }
    p = tmp_path / "sbom.json"
    p.write_text(json_module.dumps(doc))
    return str(p)


def test_build_component_graph_resolves_placeholder_chain():
    packages = [
        {"SPDXID": "SPDXRef-custom-comp", "name": "19518/helion/jre",
         "versionInfo": "1.0.0", "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-placeholder", "name": "bellsoft.zip",
         "versionInfo": "NOASSERTION", "downloadLocation": "NOASSERTION",
         "comment": "Incomplete dependency", "supplier": "Organization: Global URL"},
        {"SPDXID": "SPDXRef-real", "name": "bellsoft-jre.zip",
         "versionInfo": "abc123", "downloadLocation": "https://example.com/jre.zip",
         "supplier": "Organization: FOSSA (Global URL)"},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-custom-comp", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-placeholder"},
        {"spdxElementId": "SPDXRef-placeholder", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-custom-comp"},
        {"spdxElementId": "SPDXRef-placeholder", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-real"},
        {"spdxElementId": "SPDXRef-real", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-placeholder"},
    ]
    graph = build_component_graph(packages, relationships)
    assert graph.get("SPDXRef-real") == {"SPDXRef-custom-comp"}


def test_build_component_graph_direct_dependency_no_placeholder():
    packages = [
        {"SPDXID": "SPDXRef-comp-a", "name": "19518/helion/app-a", "versionInfo": "1.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-lodash", "name": "lodash", "versionInfo": "4.18.1",
         "supplier": "Organization: NPM"},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-lodash", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-comp-a"},
        {"spdxElementId": "SPDXRef-comp-a", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-lodash"},
    ]
    graph = build_component_graph(packages, relationships)
    assert graph.get("SPDXRef-lodash") == {"SPDXRef-comp-a"}


def test_resolve_license_uses_declared_value_when_present():
    assert resolve_license({"licenseDeclared": "Apache-2.0"}) == "Apache-2.0"


def test_resolve_license_falls_back_to_first_file_license():
    pkg = {"licenseDeclared": "NONE", "licenseInfoFromFiles": ["MIT", "BSD-3-Clause"]}
    assert resolve_license(pkg) == "MIT (derived from file scan; no declared license)"


def test_resolve_license_unknown_when_no_info_at_all():
    assert resolve_license({"licenseDeclared": "NONE"}) == "UNKNOWN - Review Required"


# ── license normalization (LicenseRef-* → SPDX) ──────────────────────────────

from sbom_lib import normalize_license, load_license_overrides
from sbom_to_oss_dependencies import flag_review_rows


def test_normalize_strips_licenseref_and_hash():
    assert normalize_license("LicenseRef-MIT-89469734") == "MIT"


def test_normalize_bare_licenseref():
    assert normalize_license("LicenseRef-MIT") == "MIT"


def test_normalize_preserves_compound_or():
    assert normalize_license("LicenseRef-MIT-91157677 OR LicenseRef-Zlib") == "MIT OR Zlib"


def test_normalize_mixed_plain_and_ref():
    got = normalize_license(
        "GPL-2.0-only WITH Classpath-exception-2.0 OR LicenseRef-EPL-2.0-69905199"
    )
    assert got == "GPL-2.0-only WITH Classpath-exception-2.0 OR EPL-2.0"


def test_normalize_embedded_with_exception():
    assert normalize_license("LicenseRef-apache-2.0-WITH-llvm-exception") == "Apache-2.0 WITH LLVM-exception"


def test_normalize_canonical_case():
    assert normalize_license("LicenseRef-apache-2.0-65368118") == "Apache-2.0"


def test_normalize_cc0():
    assert normalize_license("LicenseRef-MIT OR LicenseRef-CC0-1.0-71498429") == "MIT OR CC0-1.0"


def test_normalize_passthrough_plain_spdx():
    assert normalize_license("BSD-3-Clause") == "BSD-3-Clause"


def test_normalize_empty_and_none():
    assert normalize_license("") == ""
    assert normalize_license(None) is None


# ── curated license overrides ────────────────────────────────────────────────

def test_load_license_overrides_keys_lowercased(tmp_path):
    p = tmp_path / "lic.json"
    p.write_text(json_module.dumps({"overrides": [{"name": "FFmpeg", "license": "LGPL-2.1-or-later"}]}))
    ovr = load_license_overrides(str(p))
    assert ovr == {"ffmpeg": "LGPL-2.1-or-later"}


def test_load_license_overrides_missing_path_is_empty():
    assert load_license_overrides(None) == {}


def test_resolve_license_override_wins_over_filescan():
    pkg = {"name": "FTD2XX.Net", "licenseDeclared": "NONE",
           "licenseInfoFromFiles": ["GPL-1.0-only"]}
    assert resolve_license(pkg, {"ftd2xx.net": "LicenseRef-FTDI-FTD2XX"}) == "LicenseRef-FTDI-FTD2XX"


def test_resolve_license_override_covers_unknown():
    pkg = {"name": "org.nginx:nginx-windows", "licenseDeclared": "NONE",
           "licenseInfoFromFiles": []}
    assert resolve_license(pkg, {"org.nginx:nginx-windows": "BSD-2-Clause"}) == "BSD-2-Clause"


def test_resolve_license_declared_is_normalized():
    pkg = {"name": "left-pad", "licenseDeclared": "LicenseRef-MIT-89469734"}
    assert resolve_license(pkg, {}) == "MIT"


def test_resolve_license_filescan_fallback_normalized_with_note():
    pkg = {"name": "dcm4che", "licenseDeclared": "NONE",
           "licenseInfoFromFiles": ["LicenseRef-MPL-1.1"]}
    got = resolve_license(pkg, {})
    assert got == "MPL-1.1 (derived from file scan; no declared license)"


def test_resolve_license_still_unknown_when_nothing():
    pkg = {"name": "mystery", "licenseDeclared": "NONE", "licenseInfoFromFiles": []}
    assert resolve_license(pkg, {}) == "UNKNOWN - Review Required"


# ── review flagging ──────────────────────────────────────────────────────────

def test_flag_review_rows_detects_copyleft_unknown_and_filescan():
    rows = [
        {"name": "a", "version": "1", "license": "MIT"},
        {"name": "b", "version": "2", "license": "LGPL-3.0-only"},
        {"name": "c", "version": "3", "license": "UNKNOWN - Review Required"},
        {"name": "d", "version": "4", "license": "MPL-1.1 (derived from file scan; no declared license)"},
        {"name": "e", "version": "5", "license": "Apache-2.0"},
    ]
    flagged = flag_review_rows(rows)
    names = {f["name"] for f in flagged}
    assert names == {"b", "c", "d"}


def test_resolve_reference_prefers_homepage_then_download_location():
    assert resolve_reference({"homepage": "https://example.com"}) == "website=https://example.com"
    assert resolve_reference(
        {"homepage": "NOASSERTION", "downloadLocation": "https://dl.example.com/x.zip"}
    ) == "website=https://dl.example.com/x.zip"
    assert resolve_reference({}) == ""


def test_resolve_purl_prefers_external_ref_then_download_location():
    pkg_with_purl = {
        "externalRefs": [{"referenceType": "purl", "referenceLocator": "pkg:npm/x@1.0.0"}]
    }
    assert resolve_purl(pkg_with_purl) == "pkg:npm/x@1.0.0"
    pkg_url_only = {"downloadLocation": "https://dl.example.com/x.zip"}
    assert resolve_purl(pkg_url_only) == "https://dl.example.com/x.zip"
    assert resolve_purl({}) == ""


def test_resolve_purpose_blanks_noassertion():
    assert resolve_purpose({"summary": "NOASSERTION"}) == ""
    assert resolve_purpose({"summary": "Does things"}) == "Does things"
    assert resolve_purpose({}) == ""


def test_resolve_purpose_strips_leading_whitespace_from_summary():
    pkg = {"summary": "\n    Guava is a suite\n  "}
    assert resolve_purpose(pkg) == "Guava is a suite"


def test_resolve_purpose_strips_emoji_from_summary():
    pkg = {"summary": "👻 Primitive and flexible state management"}
    assert resolve_purpose(pkg) == "Primitive and flexible state management"


# ── purpose / reference overrides ────────────────────────────────────────────

from sbom_lib import load_package_overrides, load_purpose_overrides, load_reference_overrides


def test_load_package_overrides_returns_full_entries(tmp_path):
    p = tmp_path / "ovr.json"
    p.write_text(json_module.dumps({"overrides": [
        {"name": "Nginx", "license": "BSD-2-Clause", "purpose": "web server", "reference": "website=https://nginx.org/"}
    ]}))
    ovr = load_package_overrides(str(p))
    assert ovr["nginx"]["purpose"] == "web server"
    assert ovr["nginx"]["reference"] == "website=https://nginx.org/"


def test_load_purpose_and_reference_field_maps(tmp_path):
    p = tmp_path / "ovr.json"
    p.write_text(json_module.dumps({"overrides": [
        {"name": "a", "purpose": "does A"},
        {"name": "b", "reference": "website=https://b.example"},
        {"name": "c", "license": "MIT"}
    ]}))
    assert load_purpose_overrides(str(p)) == {"a": "does A"}
    assert load_reference_overrides(str(p)) == {"b": "website=https://b.example"}


def test_resolve_purpose_override_wins_over_empty_summary():
    pkg = {"name": "org.nginx:nginx-windows", "summary": "NOASSERTION"}
    ovr = {"org.nginx:nginx-windows": "HTTP server and reverse proxy (nginx)"}
    assert resolve_purpose(pkg, ovr) == "HTTP server and reverse proxy (nginx)"


def test_resolve_reference_override_wins_when_no_homepage():
    pkg = {"name": "cc.nssm:nssm"}
    ovr = {"cc.nssm:nssm": "website=https://nssm.cc/"}
    assert resolve_reference(pkg, ovr) == "website=https://nssm.cc/"


def test_build_rows_joins_multiple_components_for_shared_package():
    packages = [
        {"SPDXID": "SPDXRef-comp-a", "name": "19518/helion/app-a", "versionInfo": "1.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-comp-b", "name": "19518/helion/app-b", "versionInfo": "2.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-lodash", "name": "lodash", "versionInfo": "4.18.1",
         "supplier": "Organization: NPM", "summary": "Utility library",
         "licenseDeclared": "MIT", "homepage": "https://lodash.com",
         "externalRefs": [{"referenceType": "purl", "referenceLocator": "pkg:npm/lodash@4.18.1"}]},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-lodash", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-comp-a"},
        {"spdxElementId": "SPDXRef-comp-a", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-lodash"},
        {"spdxElementId": "SPDXRef-lodash", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-comp-b"},
        {"spdxElementId": "SPDXRef-comp-b", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-lodash"},
    ]
    rows, dropped, merged = build_rows(packages, relationships, {})
    assert len(rows) == 1
    assert rows[0]["name"] == "lodash"
    assert rows[0]["components"] == "19518/helion/app-a, 19518/helion/app-b"
    assert dropped == 0
    assert merged == 0


def test_build_rows_components_column_uses_canonical_names():
    from sbom_to_oss_dependencies import build_rows
    packages = [
        {"SPDXID": "SPDXRef-hel", "name": "hel-app",
         "versionInfo": "1.15.6", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-mvn", "name": "biz.videomed.tl4.installer:tl4-app-external",
         "versionInfo": "1.15.6", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-oss", "name": "left-pad", "versionInfo": "1.0.0",
         "supplier": "Organization: npm", "licenseDeclared": "MIT",
         "externalRefs": [{"referenceType": "purl",
                           "referenceLocator": "pkg:npm/left-pad@1.0.0"}]},
    ]
    # Both first-party twins depend on the OSS package.
    relationships = [
        {"spdxElementId": "SPDXRef-oss", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-hel"},
        {"spdxElementId": "SPDXRef-oss", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-mvn"},
    ]
    aliases = {"tl4-app-external": "hel-app"}
    rows, _, _ = build_rows(packages, relationships, {}, aliases=aliases)
    oss_row = next(r for r in rows if r["name"] == "left-pad")
    assert oss_row["components"] == "hel-app"


def test_build_rows_merges_genuine_duplicate_name_version():
    packages = [
        {"SPDXID": "SPDXRef-comp-a", "name": "19518/helion/app-a", "versionInfo": "1.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-comp-b", "name": "19518/helion/app-b", "versionInfo": "2.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-axios-1", "name": "axios", "versionInfo": "1.2.0",
         "supplier": "Organization: NPM", "licenseDeclared": "MIT"},
        {"SPDXID": "SPDXRef-axios-2", "name": "axios", "versionInfo": "1.2.0",
         "supplier": "Organization: NPM", "licenseDeclared": "MIT"},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-axios-1", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-comp-a"},
        {"spdxElementId": "SPDXRef-axios-2", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-comp-b"},
    ]
    rows, dropped, merged = build_rows(packages, relationships, {})
    assert len(rows) == 1
    assert rows[0]["components"] == "19518/helion/app-a, 19518/helion/app-b"
    assert merged == 1


def test_build_rows_excludes_first_party_packages():
    packages = [
        {"SPDXID": "SPDXRef-comp-a", "name": "19518/helion/app-a", "versionInfo": "1.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
    ]
    rows, dropped, merged = build_rows(packages, [], {})
    assert rows == []
    assert dropped == 0
    assert merged == 0


def test_explicit_first_party_dependency_ids_match_case_insensitively():
    packages = [
        {"SPDXID": "SPDXRef-serial", "name": "biz.videomed.tools:SerialTest",
         "versionInfo": "1.0", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-app", "name": "hel-app", "versionInfo": "2.0",
         "supplier": "Organization: Baxter"},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-serial", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-app"},
    ]
    artifact_ids = {"SERIALTEST"}

    rows, _, _ = build_rows(
        packages, relationships, {}, first_party_dependency_ids=artifact_ids
    )
    graph = build_component_graph(packages, relationships, artifact_ids)
    components = build_components(packages, {}, first_party_dependency_ids=artifact_ids)

    assert [(row["name"], row["components"]) for row in rows] == [
        ("biz.videomed.tools:SerialTest", "hel-app")
    ]
    assert graph == {"SPDXRef-serial": {"SPDXRef-app"}}
    assert [component["name"] for component in components] == ["hel-app"]


def test_build_rows_drops_placeholders_and_counts_them():
    packages = [
        {"SPDXID": "SPDXRef-placeholder", "name": "x.zip", "versionInfo": "NOASSERTION",
         "downloadLocation": "NOASSERTION", "comment": "Incomplete dependency",
         "supplier": "Organization: Global URL"},
        {"SPDXID": "SPDXRef-real", "name": "x.zip", "versionInfo": "1.0",
         "downloadLocation": "https://example.com/x.zip",
         "supplier": "Organization: FOSSA (Global URL)", "licenseDeclared": "MIT"},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-placeholder", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-real"},
        {"spdxElementId": "SPDXRef-real", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-placeholder"},
    ]
    rows, dropped, merged = build_rows(packages, relationships, {})
    assert dropped == 1
    assert len(rows) == 1
    assert rows[0]["name"] == "x.zip"


def test_build_rows_sorts_by_name_then_version():
    packages = [
        {"SPDXID": "SPDXRef-b", "name": "banana", "versionInfo": "1.0", "supplier": "Organization: NPM"},
        {"SPDXID": "SPDXRef-a", "name": "apple", "versionInfo": "2.0", "supplier": "Organization: NPM"},
    ]
    rows, _, _ = build_rows(packages, [], {})
    assert [r["name"] for r in rows] == ["apple", "banana"]


def test_build_components_returns_first_party_packages_sorted_by_name():
    packages = [
        {"SPDXID": "SPDXRef-comp-b", "name": "19518/helion/zeta", "versionInfo": "2.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-comp-a", "name": "19518/helion/alpha", "versionInfo": "1.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-lodash", "name": "lodash", "versionInfo": "4.18.1",
         "supplier": "Organization: NPM"},
    ]
    components = build_components(packages, {})
    assert [c["name"] for c in components] == ["19518/helion/alpha", "19518/helion/zeta"]
    assert components[0]["version"] == "1.0"


def test_build_components_collapses_aliased_first_party_twins():
    from sbom_to_oss_dependencies import build_components
    packages = [
        {"SPDXID": "SPDXRef-hel", "name": "hel-app",
         "versionInfo": "1.15.6", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-mvn", "name": "biz.videomed.tl4.installer:tl4-app-external",
         "versionInfo": "1.15.6", "supplier": "Organization: Maven"},
    ]
    aliases = {"tl4-app-external": "hel-app"}
    components = build_components(packages, {}, aliases)
    names = [c["name"] for c in components]
    assert names == ["hel-app"]
    assert len(components) == 1


def test_build_components_collapses_same_artifact_twins_without_alias():
    """Pairs sharing an artifact id/version (e.g. nms-gateway-helper) must
    collapse even when not in the curated alias map, preferring the 'hel-'
    display name."""
    from sbom_to_oss_dependencies import build_components
    packages = [
        {"SPDXID": "SPDXRef-mvn", "name": "biz.videomed.tl4.installer:nms-gateway-helper",
         "versionInfo": "1.5.5", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-hel", "name": "hel-nms-gateway-helper",
         "versionInfo": "1.5.5", "supplier": "Organization: Baxter"},
    ]
    components = build_components(packages, {})
    assert [c["name"] for c in components] == ["hel-nms-gateway-helper"]


def test_build_components_no_aliases_keeps_both():
    from sbom_to_oss_dependencies import build_components
    packages = [
        {"SPDXID": "SPDXRef-hel", "name": "hel-app",
         "versionInfo": "1.15.6", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-mvn", "name": "biz.videomed.tl4.installer:tl4-app-external",
         "versionInfo": "1.15.6", "supplier": "Organization: Maven"},
    ]
    components = build_components(packages, {})
    assert len(components) == 2


def test_build_excel_writes_expected_sheet_dimensions_and_content(tmp_path):
    packages = [
        {"SPDXID": "SPDXRef-comp-a", "name": "19518/helion/app-a", "versionInfo": "1.0",
         "supplier": "Organization: FOSSA (Custom (provided build))"},
        {"SPDXID": "SPDXRef-lodash", "name": "lodash", "versionInfo": "4.18.1",
         "supplier": "Organization: NPM", "licenseDeclared": "MIT",
         "homepage": "https://lodash.com",
         "externalRefs": [{"referenceType": "purl", "referenceLocator": "pkg:npm/lodash@4.18.1"}]},
    ]
    relationships = [
        {"spdxElementId": "SPDXRef-lodash", "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-comp-a"},
        {"spdxElementId": "SPDXRef-comp-a", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": "SPDXRef-lodash"},
    ]
    sbom = _write_sbom(tmp_path, packages, relationships)
    out = str(tmp_path / "out.xlsx")

    build_excel(sbom_path=sbom, template_path=OSS_TEMPLATE, output_path=out)

    wb = _openpyxl.load_workbook(out)
    ws = wb["Dependencies (OTS SOUP)"]
    assert ws.dimensions == "A1:J2"
    assert ws.cell(2, 1).value == "lodash"
    assert ws.cell(2, 3).value == "4.18.1"
    assert ws.cell(2, 7).value == "MIT"
    assert ws.cell(2, 8).value == "website=https://lodash.com"
    assert ws.cell(2, 9).value == "19518/helion/app-a"
    assert ws.cell(2, 10).value == "pkg:npm/lodash@4.18.1"

    ws2 = wb["SW-SYS Components (Ref-only)"]
    assert ws2.dimensions == "A1:C3"
    assert ws2.cell(3, 1).value == "19518/helion/app-a"
    assert ws2.cell(3, 2).value == "1.0"
    assert "Test Product" in ws2.cell(1, 1).value


def test_build_excel_classifies_configured_first_party_artifacts_as_dependencies(tmp_path):
    dependency_names = [
        "biz.videomed.t4.tools:lang-manifest",
        "biz.videomed.t4.tools:LDBootloader",
        "biz.videomed.t4.tools:SerialTest",
        "biz.videomed.t4.tools:SystemTest",
    ]
    packages = [
        {"SPDXID": "SPDXRef-lang-manifest", "name": dependency_names[0],
         "versionInfo": "1.0.1", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-ldbootloader", "name": dependency_names[1],
         "versionInfo": "2.0.2", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-serialtest", "name": dependency_names[2],
         "versionInfo": "3.0.3", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-systemtest", "name": dependency_names[3],
         "versionInfo": "4.0.4", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-hel-app", "name": "hel-app",
         "versionInfo": "5.0.5", "supplier": "Organization: Baxter"},
    ]
    relationships = [
        {"spdxElementId": package["SPDXID"], "relationshipType": "DEPENDENCY_OF",
         "relatedSpdxElement": "SPDXRef-hel-app"}
        for package in packages[:2]
    ] + [
        {"spdxElementId": "SPDXRef-hel-app", "relationshipType": "DEPENDS_ON",
         "relatedSpdxElement": package["SPDXID"]}
        for package in packages[2:4]
    ]

    sbom = _write_sbom(tmp_path, packages, relationships)
    with open(sbom, encoding="utf-8") as f:
        source_document_before = json_module.load(f)

    out = tmp_path / "out.xlsx"
    build_excel(sbom, OSS_TEMPLATE, str(out))

    wb = _openpyxl.load_workbook(str(out))
    dependencies_ws = wb["Dependencies (OTS SOUP)"]
    dependency_rows = [
        (
            dependencies_ws.cell(row, 1).value,
            dependencies_ws.cell(row, 3).value,
            dependencies_ws.cell(row, 9).value,
        )
        for row in range(2, dependencies_ws.max_row + 1)
    ]
    assert dependency_rows == [
        (dependency_names[0], "1.0.1", "hel-app"),
        (dependency_names[1], "2.0.2", "hel-app"),
        (dependency_names[2], "3.0.3", "hel-app"),
        (dependency_names[3], "4.0.4", "hel-app"),
    ]

    components_ws = wb["SW-SYS Components (Ref-only)"]
    component_names = [
        components_ws.cell(row, 1).value
        for row in range(3, components_ws.max_row + 1)
        if components_ws.cell(row, 1).value
    ]
    assert component_names == ["hel-app"]

    with open(sbom, encoding="utf-8") as f:
        source_document_after = json_module.load(f)
    assert source_document_after == source_document_before


def test_build_excel_end_to_end_collapses_twins_in_components_sheet(tmp_path):
    from sbom_to_oss_dependencies import build_excel
    packages = [
        {"SPDXID": "SPDXRef-hel", "name": "hel-app",
         "versionInfo": "1.15.6", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-mvn", "name": "biz.videomed.tl4.installer:tl4-app-external",
         "versionInfo": "1.15.6", "supplier": "Organization: Maven"},
    ]
    sbom = _write_sbom(tmp_path, packages)
    aliases_path = tmp_path / "aliases.json"
    aliases_path.write_text(json_module.dumps({"tl4-app-external": "hel-app"}))
    out = tmp_path / "out.xlsx"
    build_excel(sbom, OSS_TEMPLATE, str(out), component_aliases_path=str(aliases_path))

    wb = _openpyxl.load_workbook(str(out))
    ws = wb["SW-SYS Components (Ref-only)"]
    names = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)
             if ws.cell(r, 1).value]
    assert names == ["hel-app"]


def test_build_excel_drops_explicitly_excluded_components(tmp_path):
    from sbom_to_oss_dependencies import build_excel
    packages = [
        {"SPDXID": "SPDXRef-hel", "name": "hel-app",
         "versionInfo": "1.15.6", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-lic", "name": "biz.videomed:licensing",
         "versionInfo": "0.6.8", "supplier": "Organization: Maven"},
    ]
    sbom = _write_sbom(tmp_path, packages)
    excluded_path = tmp_path / "excluded.json"
    excluded_path.write_text(json_module.dumps({"licensing": "test"}))
    out = tmp_path / "out.xlsx"
    build_excel(sbom, OSS_TEMPLATE, str(out), excluded_components_path=str(excluded_path))

    wb = _openpyxl.load_workbook(str(out))
    ws = wb["SW-SYS Components (Ref-only)"]
    names = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)
             if ws.cell(r, 1).value]
    assert names == ["hel-app"]


def test_build_excel_applies_oss_deps_only_exclusions_by_default(tmp_path):
    packages = [
        {"SPDXID": "SPDXRef-invalidated", "name": "hel-invalidated-tokens-api",
         "versionInfo": "1.0", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-mvn-invalidated",
         "name": "biz.videomed.tl4:invalidated-tokens-api",
         "versionInfo": "1.0", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-broker", "name": "hel-nms-broker-api",
         "versionInfo": "1.0", "supplier": "Organization: Baxter"},
        {"SPDXID": "SPDXRef-mvn-broker",
         "name": "biz.videomed.nexxis:nms-broker-api",
         "versionInfo": "1.0", "supplier": "Organization: Maven"},
        {"SPDXID": "SPDXRef-app", "name": "hel-app",
         "versionInfo": "1.0", "supplier": "Organization: Baxter"},
    ]
    sbom = _write_sbom(tmp_path, packages)
    out = tmp_path / "out.xlsx"

    build_excel(sbom, OSS_TEMPLATE, str(out))

    with open(sbom, encoding="utf-8") as f:
        source_names = [package["name"] for package in json_module.load(f)["packages"]]
    assert source_names == [
        "hel-invalidated-tokens-api",
        "biz.videomed.tl4:invalidated-tokens-api",
        "hel-nms-broker-api",
        "biz.videomed.nexxis:nms-broker-api",
        "hel-app",
    ]

    wb = _openpyxl.load_workbook(str(out))
    ws = wb["SW-SYS Components (Ref-only)"]
    names = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)
             if ws.cell(r, 1).value]
    assert names == ["hel-app"]
