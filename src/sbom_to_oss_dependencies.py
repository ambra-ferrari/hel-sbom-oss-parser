#!/usr/bin/env python3
"""
sbom_to_oss_dependencies.py — Populate the OTS-SOUP OSS dependencies Excel
template (`oss-template.xlsx`, Helion product line) from an SPDX 2.3 JSON
SBOM file.

Two sheets are produced:
  - "Dependencies (OTS SOUP)": one row per external OSS package, with the
    internal Helion components that depend on it listed in a single
    comma-joined cell.
  - "SW-SYS Components (Ref-only)": the internal (first-party) Helion
    components themselves, for reference.

See docs/superpowers/specs/2026-09-29-oss-dependencies-agent-design.md for
the full design rationale.
"""

import argparse
from importlib import metadata
import json
import re
import sys
import sysconfig
from pathlib import Path

try:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
except ImportError:
    sys.exit("Missing dependency: pip install openpyxl")

from sbom_lib import (
    is_incomplete_dependency,
    is_first_party,
    is_excluded_component,
    resolve_version,
    load_version_overrides,
    canonical_component_name,
    first_party_artifact_key,
    load_component_aliases,
    load_excluded_components,
    load_license_overrides,
    load_purpose_overrides,
    load_ref_overrides,
    load_reference_overrides,
    load_vendor_overrides,
    normalize_license,
    normalize_purpose_text,
    copy_row_style,
)

DEFAULT_OSS_DEPS_EXCLUDED_COMPONENTS = (
    Path(__file__).resolve().parent.parent
    / "config"
    / "oss_deps_excluded_components.json"
)
DEFAULT_OSS_DEPS_FIRST_PARTY_DEPENDENCIES = (
    Path(__file__).resolve().parent.parent
    / "config"
    / "oss_deps_first_party_dependencies.json"
)


def resolve_first_party_dependency_config_path(
        config_path: str | Path | None = None) -> Path:
    if config_path is not None:
        path = Path(config_path)
        if path.is_file():
            return path
        raise FileNotFoundError(f"First-party dependency config file not found: {path}")

    if DEFAULT_OSS_DEPS_FIRST_PARTY_DEPENDENCIES.is_file():
        return DEFAULT_OSS_DEPS_FIRST_PARTY_DEPENDENCIES

    try:
        distribution = metadata.distribution("hel-sbom-oss-parser")
    except metadata.PackageNotFoundError:
        distribution = None
    if distribution is not None:
        expected_parts = (
            "share", "hel-sbom-oss-parser",
            "oss_deps_first_party_dependencies.json",
        )
        for installed_file in distribution.files or ():
            if tuple(installed_file.parts[-3:]) != expected_parts:
                continue
            located_path = Path(distribution.locate_file(installed_file))
            if located_path.is_file():
                return located_path

    packaged_path = (
        Path(sysconfig.get_path("data"))
        / "share"
        / "hel-sbom-oss-parser"
        / "oss_deps_first_party_dependencies.json"
    )
    if packaged_path.is_file():
        return packaged_path

    raise FileNotFoundError(
        "First-party dependency config file not found in source tree, "
        f"installed distribution, or installation data directory: {packaged_path}"
    )


def load_first_party_dependency_ids(config_path: str | Path) -> set[str]:
    path = Path(config_path)
    try:
        with path.open(encoding="utf-8") as config_file:
            config = json.load(config_file)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"First-party dependency config is not valid JSON: {path} ({exc})"
        ) from exc
    if not isinstance(config, dict):
        raise ValueError(
            f"First-party dependency config must contain a JSON object: {path}"
        )
    return {
        str(artifact_id).strip().casefold()
        for artifact_id in config
        if str(artifact_id).strip()
    }


# ── component graph ───────────────────────────────────────────────────────────

def is_configured_first_party_dependency(pkg: dict, artifact_ids: set[str]) -> bool:
    name = (pkg.get("name") or "").strip()
    artifact_id = name.rsplit(":", 1)[-1].strip().casefold()
    return artifact_id in artifact_ids


def _normalize_first_party_dependency_ids(artifact_ids: set[str] | None) -> set[str]:
    return {artifact_id.strip().casefold() for artifact_id in artifact_ids or set()}


def build_component_graph(
        packages: list[dict],
        relationships: list[dict],
        first_party_dependency_ids: set[str] | None = None) -> dict[str, set[str]]:
    """Map every real package's SPDXID to the set of internal-component
    SPDXIDs that depend on it.

    Direct edges are accepted in either SPDX encoding:
      1. `DEPENDENCY_OF`: package subject -> component object.
      2. `DEPENDS_ON`: component subject -> package object.
    FOSSA sometimes routes a component's dependency through an "Incomplete
    dependency" placeholder package (component -> placeholder -> real). In
    that case, the placeholder's resolved components are attributed directly
    to the real package it `DEPENDS_ON`.
    """
    first_party_dependency_ids = _normalize_first_party_dependency_ids(
        first_party_dependency_ids
    )
    component_ids = {
        pkg["SPDXID"] for pkg in packages
        if is_first_party(pkg)
        and not is_configured_first_party_dependency(pkg, first_party_dependency_ids)
    }
    package_ids = {pkg["SPDXID"] for pkg in packages}
    placeholder_ids = {pkg["SPDXID"] for pkg in packages if is_incomplete_dependency(pkg)}

    pkg_to_components: dict[str, set[str]] = {}
    for rel in relationships:
        relationship_type = rel.get("relationshipType")
        source = rel.get("spdxElementId")
        related = rel.get("relatedSpdxElement")
        if relationship_type == "DEPENDENCY_OF" and related in component_ids:
            pkg_to_components.setdefault(source, set()).add(related)
        elif (relationship_type == "DEPENDS_ON"
              and source in component_ids and related in package_ids):
            pkg_to_components.setdefault(related, set()).add(source)

    for rel in relationships:
        if rel.get("relationshipType") != "DEPENDS_ON":
            continue
        source = rel.get("spdxElementId")
        if source in placeholder_ids:
            comps = pkg_to_components.get(source)
            if comps:
                real_id = rel.get("relatedSpdxElement")
                pkg_to_components.setdefault(real_id, set()).update(comps)

    return pkg_to_components


# ── field resolution ──────────────────────────────────────────────────────────

def resolve_license(pkg: dict, license_overrides: dict | None = None) -> str:
    """Curated override (by package name) wins; else declared license; else the
    first file-scanned license with a note; else an explicit review flag. All
    non-override results are normalized from FOSSA LicenseRef-* to clean SPDX."""
    license_overrides = license_overrides or {}
    name = (pkg.get("name") or "").strip().lower()
    if name in license_overrides:
        return license_overrides[name]
    declared = pkg.get("licenseDeclared")
    if declared and declared != "NONE":
        return normalize_license(declared)
    files = pkg.get("licenseInfoFromFiles") or []
    if files:
        return f"{normalize_license(files[0])} (derived from file scan; no declared license)"
    return "UNKNOWN - Review Required"


def resolve_reference(pkg: dict, reference_overrides: dict | None = None) -> str:
    """Curated override (by package name) wins; else `website=<url>` using
    homepage, falling back to the download location."""
    reference_overrides = reference_overrides or {}
    name = (pkg.get("name") or "").strip().lower()
    if name in reference_overrides:
        return reference_overrides[name]
    homepage = pkg.get("homepage")
    if homepage and homepage != "NOASSERTION":
        return f"website={homepage}"
    download = pkg.get("downloadLocation")
    if download and download != "NOASSERTION":
        return f"website={download}"
    return ""


def _resolve_purl_with_priority(
        pkg: dict, ref_overrides: dict[str, str] | None = None) -> tuple[str, int]:
    """Return a package reference and its precedence: PURL, download, then override."""
    for ref in pkg.get("externalRefs") or []:
        if ref.get("referenceType") == "purl":
            locator = (ref.get("referenceLocator") or "").strip()
            if locator and locator != "NOASSERTION":
                return locator, 2
    download = (pkg.get("downloadLocation") or "").strip()
    if download and download != "NOASSERTION":
        return download, 1
    ref_overrides = ref_overrides or {}
    name = (pkg.get("name") or "").strip().lower()
    return ref_overrides.get(name, ""), 0


def resolve_purl(pkg: dict, ref_overrides: dict[str, str] | None = None) -> str:
    """Resolve a PURL/download URL, then use an explicit per-package Ref fallback."""
    return _resolve_purl_with_priority(pkg, ref_overrides)[0]


def resolve_purpose(pkg: dict, purpose_overrides: dict | None = None) -> str:
    """Curated override (by package name) wins; else the package summary,
    blanked when SPDX has no real value for it. The result is normalized
    (see `normalize_purpose_text`) to strip stray leading/trailing
    whitespace and decorative emoji some upstream summaries carry."""
    purpose_overrides = purpose_overrides or {}
    name = (pkg.get("name") or "").strip().lower()
    if name in purpose_overrides:
        return normalize_purpose_text(purpose_overrides[name])
    summary = pkg.get("summary")
    if not summary or summary == "NOASSERTION":
        return ""
    return normalize_purpose_text(summary)


# ── row assembly ──────────────────────────────────────────────────────────────

def build_rows(packages: list[dict], relationships: list[dict], version_overrides: dict,
               license_overrides: dict | None = None,
               purpose_overrides: dict | None = None,
               reference_overrides: dict | None = None,
               aliases: dict[str, str] | None = None,
               first_party_dependency_ids: set[str] | None = None,
               vendor_overrides: dict[str, str] | None = None,
               ref_overrides: dict[str, str] | None = None,
               component_packages: list[dict] | None = None):
    """Build one output row per unique OSS package (name + version).

    Returns (rows, placeholders_dropped, duplicates_merged):
      - rows: list of dicts with keys name/version/purpose/license/reference/
        components (comma-joined str)/ref, sorted by (name.lower(), version).
      - placeholders_dropped: count of FOSSA "Incomplete dependency" entries excluded.
      - duplicates_merged: count of packages that shared a (name, version) key
        with an already-written row and had their components merged into it
        instead of creating a duplicate row (happens with aggregated
        multi-project SBOMs).
    """
    first_party_dependency_ids = _normalize_first_party_dependency_ids(
        first_party_dependency_ids
    )
    graph_packages = component_packages if component_packages is not None else packages
    pkg_to_components = build_component_graph(
        graph_packages, relationships, first_party_dependency_ids
    )
    pkgs_by_id = {pkg["SPDXID"]: pkg for pkg in graph_packages}
    aliases = aliases or {}
    vendor_overrides = vendor_overrides or {}
    ref_overrides = ref_overrides or {}

    placeholders_dropped = sum(1 for pkg in packages if is_incomplete_dependency(pkg))

    oss_packages = [
        pkg for pkg in packages
        if (not is_first_party(pkg)
            or is_configured_first_party_dependency(pkg, first_party_dependency_ids))
        and not is_incomplete_dependency(pkg)
    ]

    rows_by_key: dict[tuple[str, str], dict] = {}
    ref_priorities: dict[tuple[str, str], int] = {}
    order: list[tuple[str, str]] = []
    duplicates_merged = 0

    for pkg in oss_packages:
        name = (pkg.get("name") or "").strip() or resolve_purl(pkg) or pkg.get("downloadLocation", "")
        version = resolve_version(pkg, version_overrides)
        key = (name.lower(), version)

        component_ids = pkg_to_components.get(pkg["SPDXID"], set())
        component_names = {
            canonical_component_name(pkgs_by_id[c]["name"], aliases)
            for c in component_ids if c in pkgs_by_id
        }

        if key in rows_by_key:
            rows_by_key[key]["components"].update(component_names)
            ref, ref_priority = _resolve_purl_with_priority(pkg, ref_overrides)
            if ref_priority > ref_priorities[key]:
                rows_by_key[key]["ref"] = ref
                ref_priorities[key] = ref_priority
            duplicates_merged += 1
            continue

        purpose = resolve_purpose(pkg, purpose_overrides)
        if purpose:
            purpose = purpose[0].upper() + purpose[1:]
        ref, ref_priority = _resolve_purl_with_priority(pkg, ref_overrides)

        rows_by_key[key] = {
            "name": name,
            "version": version,
            "vendor": vendor_overrides.get(
                (pkg.get("name") or "").strip().lower(), "Open Source"
            ),
            "purpose": purpose,
            "license": resolve_license(pkg, license_overrides),
            "reference": resolve_reference(pkg, reference_overrides),
            "components": set(component_names),
            "ref": ref,
        }
        ref_priorities[key] = ref_priority
        order.append(key)

    rows = [rows_by_key[key] for key in order]
    rows.sort(key=lambda r: (r["name"].lower(), r["version"]))
    for row in rows:
        row["components"] = ", ".join(sorted(row["components"], key=str.lower))

    return rows, placeholders_dropped, duplicates_merged


_COPYLEFT_RE = re.compile(
    r"\b(GPL|LGPL|AGPL|MPL|CDDL|EPL|SSPL|OSL|CC-BY|CC0|OFL|MS-PL|MS-NET|Zlib)\b",
    re.IGNORECASE,
)


def flag_review_rows(rows: list[dict]) -> list[dict]:
    """Rows whose license needs human/legal review: copyleft or notable family,
    an explicit UNKNOWN flag, or a low-confidence file-scan derivation."""
    flagged = []
    for row in rows:
        lic = row["license"]
        if (lic == "UNKNOWN - Review Required"
                or "(derived from file scan" in lic
                or _COPYLEFT_RE.search(lic)):
            flagged.append(row)
    return flagged


def build_components(packages: list[dict], version_overrides: dict,
                     aliases: dict[str, str] | None = None,
                     first_party_dependency_ids: set[str] | None = None) -> list[dict]:
    """Internal Helion components for the 'SW-SYS Components (Ref-only)' sheet.

    Names are resolved to their canonical form via the curated alias map, then
    grouped by (first_party_artifact_key, version) so the two representations
    of the same first-party component collapse to one row even when they
    aren't in the alias map (e.g. 'biz.videomed…:invalidated-tokens-api' and
    'hel-invalidated-tokens-api' share the same artifact id and version).
    Within a group, the 'hel-' name is preferred for display. Sorted by name.
    """
    aliases = aliases or {}
    first_party_dependency_ids = _normalize_first_party_dependency_ids(
        first_party_dependency_ids
    )
    names_by_key: dict[tuple[str, str], list[str]] = {}
    for pkg in packages:
        if (not is_first_party(pkg)
                or is_configured_first_party_dependency(pkg, first_party_dependency_ids)):
            continue
        name = canonical_component_name(pkg.get("name", ""), aliases)
        version = resolve_version(pkg, version_overrides)
        key = (first_party_artifact_key(name), version)
        names_by_key.setdefault(key, []).append(name)

    def preferred_name(names: list[str]) -> str:
        for n in names:
            if n.lower().startswith("hel-"):
                return n
        return names[0]

    components = [
        {"name": preferred_name(names), "version": key[1]}
        for key, names in names_by_key.items()
    ]
    components.sort(key=lambda c: c["name"].lower())
    return components


# ── sheet writing ─────────────────────────────────────────────────────────────

def _unmerge_all(ws):
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def _clear_row_values(ws, start_row: int, end_row: int, num_cols: int):
    for r in range(start_row, end_row + 1):
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def write_oss_dependencies_sheet(wb, rows: list[dict]):
    """Write the 'Dependencies (OTS SOUP)' sheet, trimmed to exactly
    1 header row + len(rows) data rows and 10 columns (A-J)."""
    ws = wb["Dependencies (OTS SOUP)"]
    num_cols = 10  # A..J
    _unmerge_all(ws)
    _clear_row_values(ws, 2, ws.max_row, num_cols)

    for i, row in enumerate(rows):
        r = i + 2
        if r > 2:
            copy_row_style(ws, 2, r, num_cols)
        ws.cell(r, 1).value = row["name"]
        ws.cell(r, 2).value = "See Website"
        ws.cell(r, 3).value = row["version"]
        ws.cell(r, 4).value = row["vendor"]
        ws.cell(r, 5).value = row["purpose"]
        ws.cell(r, 6).value = "Same as Purpose"
        ws.cell(r, 7).value = row["license"]
        ws.cell(r, 8).value = row["reference"]
        ws.cell(r, 9).value = row["components"]
        ws.cell(r, 10).value = row["ref"]

    last_row = len(rows) + 1
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)
    if ws.max_column > num_cols:
        ws.delete_cols(num_cols + 1, ws.max_column - num_cols)

    ws.auto_filter.ref = f"A1:J{last_row}"
    ws.freeze_panes = "A2"


def write_components_sheet(wb, components: list[dict], product_name: str):
    """Write the 'SW-SYS Components (Ref-only)' sheet, trimmed to exactly
    2 header rows + len(components) data rows and 3 columns (A-C)."""
    ws = wb["SW-SYS Components (Ref-only)"]
    num_cols = 3  # A..C
    _unmerge_all(ws)
    _clear_row_values(ws, 3, ws.max_row, num_cols)

    ws.cell(1, 1).value = (
        f"List of components of {product_name} developed internally "
        "(not SOUP/OTS) - for reference only"
    )

    for i, comp in enumerate(components):
        r = i + 3
        if r > 3:
            copy_row_style(ws, 3, r, num_cols)
        ws.cell(r, 1).value = comp["name"]
        ws.cell(r, 2).value = comp["version"]
        ws.cell(r, 3).value = ""

    last_row = len(components) + 2
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)
    if ws.max_column > num_cols:
        ws.delete_cols(num_cols + 1, ws.max_column - num_cols)


# ── main ──────────────────────────────────────────────────────────────────────

def build_excel(sbom_path: str, template_path: str, output_path: str,
                 product_name: str | None = None,
                 version_overrides_path: str | None = None,
                 license_overrides_path: str | None = None,
                 component_aliases_path: str | None = None,
                 excluded_components_path: str | None = None,
                 oss_deps_excluded_components_path: str | None = None,
                 first_party_dependency_ids: set[str] | None = None,
                 first_party_dependency_config_path: str | None = None):
    version_overrides = load_version_overrides(version_overrides_path)
    license_overrides = load_license_overrides(license_overrides_path)
    purpose_overrides = load_purpose_overrides(license_overrides_path)
    ref_overrides = load_ref_overrides(license_overrides_path)
    reference_overrides = load_reference_overrides(license_overrides_path)
    vendor_overrides = load_vendor_overrides(license_overrides_path)
    aliases = load_component_aliases(component_aliases_path)
    excluded = load_excluded_components(excluded_components_path)
    oss_deps_excluded_path = (
        oss_deps_excluded_components_path
        if oss_deps_excluded_components_path is not None
        else str(DEFAULT_OSS_DEPS_EXCLUDED_COMPONENTS)
    )
    excluded.update(load_excluded_components(oss_deps_excluded_path))
    if first_party_dependency_ids is None:
        first_party_dependency_path = resolve_first_party_dependency_config_path(
            first_party_dependency_config_path
        )
        first_party_dependency_ids = load_first_party_dependency_ids(
            first_party_dependency_path
        )
    elif first_party_dependency_config_path is not None:
        raise ValueError(
            "first_party_dependency_config_path cannot be combined with "
            "first_party_dependency_ids"
        )
    first_party_dependency_ids = _normalize_first_party_dependency_ids(
        first_party_dependency_ids
    )

    with open(sbom_path, encoding="utf-8") as f:
        sbom = json.load(f)

    all_packages = sbom.get("packages", [])
    packages = [
        pkg for pkg in all_packages
        if not is_excluded_component(pkg.get("name", ""), excluded)
    ]
    relationships = sbom.get("relationships", [])
    resolved_name = product_name or sbom.get("name", "Unknown Product")

    rows, placeholders_dropped, duplicates_merged = build_rows(
        packages, relationships, version_overrides,
        license_overrides, purpose_overrides, reference_overrides,
        aliases=aliases,
        first_party_dependency_ids=first_party_dependency_ids,
        vendor_overrides=vendor_overrides,
        ref_overrides=ref_overrides,
        component_packages=all_packages,
    )
    components = build_components(
        packages, version_overrides, aliases,
        first_party_dependency_ids=first_party_dependency_ids,
    )

    unknown_licenses = sum(1 for row in rows if row["license"] == "UNKNOWN - Review Required")
    fallback_licenses = sum(1 for row in rows if "(derived from file scan" in row["license"])
    review_rows = flag_review_rows(rows)

    wb = openpyxl.load_workbook(template_path)
    write_oss_dependencies_sheet(wb, rows)
    write_components_sheet(wb, components, resolved_name)
    wb.save(output_path)

    print(
        f"✅  Written {len(rows)} OSS rows + {len(components)} components → {output_path}\n"
        f"    placeholders dropped: {placeholders_dropped}, "
        f"duplicate rows merged: {duplicates_merged}, "
        f"licenses from file-scan fallback: {fallback_licenses}, "
        f"licenses UNKNOWN: {unknown_licenses}"
    )
    if review_rows:
        print(f"\n⚠️  {len(review_rows)} rows need license review (copyleft / UNKNOWN / file-scan):")
        for r in review_rows:
            print(f"    - {r['name']} {r['version']}: {r['license']}")


def main():
    parser = argparse.ArgumentParser(
        description="Populate the OTS-SOUP OSS dependencies xlsx from an SPDX 2.3 JSON SBOM."
    )
    parser.add_argument("--sbom", required=True, help="Path to SPDX 2.3 JSON SBOM file")
    parser.add_argument("--template", required=True, help="Path to the OTS-SOUP Excel template (.xlsx)")
    parser.add_argument("--output", required=True, help="Path for the output Excel file (.xlsx)")
    parser.add_argument(
        "--product-name", default=None,
        help="Product name shown in the SW-SYS Components sheet title "
             "(defaults to the SBOM document's own 'name' field)"
    )
    parser.add_argument(
        "--version-overrides", default=None,
        help="Path to a JSON file of {package name or SPDXID: real version} "
             "to use instead of the SPDX versionInfo."
    )
    parser.add_argument(
        "--license-overrides", default=None,
        help="Path to a JSON file of curated {package name: SPDX license} overrides "
             "(e.g. license_overrides.json) for external binaries FOSSA can't resolve."
    )
    parser.add_argument(
        "--component-aliases", default=None,
        help="Path to a JSON file of {alternate first-party name: canonical "
             "hel-* name} (e.g. component_aliases.json) to collapse the two "
             "representations of the same internal component into one row."
    )
    parser.add_argument(
        "--excluded-components", default=None,
        help="Path to a JSON file of {short artifact name: reason} (e.g. "
             "excluded_components.json) for components that must never "
             "appear in the output, regardless of party/version."
    )
    parser.add_argument(
        "--oss-deps-excluded-components",
        default=str(DEFAULT_OSS_DEPS_EXCLUDED_COMPONENTS),
        help="Path to a JSON file of components excluded only from this OSS "
             "dependencies deliverable (defaults to config/"
             "oss_deps_excluded_components.json).",
    )
    parser.add_argument(
        "--oss-deps-first-party-dependencies",
        default=None,
        help="Path to a JSON file of first-party artifact IDs treated as "
             "dependencies only in the OSS-dependencies workbook (defaults to "
             "the source-tree config or packaged config).",
    )
    args = parser.parse_args()

    build_excel(
        sbom_path=args.sbom,
        template_path=args.template,
        output_path=args.output,
        product_name=args.product_name,
        version_overrides_path=args.version_overrides,
        license_overrides_path=args.license_overrides,
        component_aliases_path=args.component_aliases,
        excluded_components_path=args.excluded_components,
        oss_deps_excluded_components_path=args.oss_deps_excluded_components,
        first_party_dependency_config_path=args.oss_deps_first_party_dependencies,
    )


if __name__ == "__main__":
    main()
